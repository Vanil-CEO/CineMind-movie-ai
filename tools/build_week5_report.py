from __future__ import annotations

from pathlib import Path
import textwrap

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "docs" / "reports" / "week-5"
SCREEN_DIR = REPORT_DIR / "screenshots"
OUT_DOCX = REPORT_DIR / "CineMind_звіт_тиждень_5.docx"
USE_CASE = REPORT_DIR / "use_case_week5.png"


def font(size: int, bold: bool = False, mono: bool = False) -> ImageFont.FreeTypeFont:
    candidates = []
    if mono:
        candidates += [
            "C:/Windows/Fonts/consola.ttf",
            "C:/Windows/Fonts/cour.ttf",
        ]
    if bold:
        candidates += [
            "C:/Windows/Fonts/arialbd.ttf",
            "C:/Windows/Fonts/calibrib.ttf",
        ]
    candidates += [
        "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/calibri.ttf",
        "C:/Windows/Fonts/segoeui.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size=size)
    return ImageFont.load_default()


def rounded(draw: ImageDraw.ImageDraw, box, radius, fill, outline=None, width=1):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def make_use_case_diagram() -> None:
    img = Image.new("RGB", (1600, 950), "#F6F8FB")
    d = ImageDraw.Draw(img)
    title_f = font(44, bold=True)
    label_f = font(27, bold=True)
    small_f = font(22)

    d.text((60, 42), "Діаграма прецедентів CineMind, тиждень 5", font=title_f, fill="#111827")
    d.text((60, 96), "Користувацькі сценарії після редизайну, AI Agent Studio, FastAPI та PostgreSQL", font=small_f, fill="#475569")

    system = (350, 150, 1260, 860)
    rounded(d, system, 30, "#FFFFFF", "#CBD5E1", 3)
    d.text((690, 176), "CineMind", font=label_f, fill="#0F172A")

    actors = {
        "Користувач": (120, 320),
        "Адміністратор": (1360, 320),
        "AI Agent": (1360, 650),
    }
    for name, (x, y) in actors.items():
        d.ellipse((x - 28, y - 78, x + 28, y - 22), fill="#0A84FF")
        d.line((x, y - 22, x, y + 78), fill="#111827", width=5)
        d.line((x - 52, y + 10, x + 52, y + 10), fill="#111827", width=5)
        d.line((x, y + 78, x - 48, y + 140), fill="#111827", width=5)
        d.line((x, y + 78, x + 48, y + 140), fill="#111827", width=5)
        d.text((x - 84, y + 160), name, font=small_f, fill="#111827")

    cases = [
        ("Реєстрація / вхід", (520, 255)),
        ("Вибір кіновподобань", (830, 255)),
        ("Перегляд персональної стрічки", (565, 395)),
        ("Перегляд трейлера", (880, 395)),
        ("Оцінювання зірочками", (555, 535)),
        ("Коментарі до фільмів", (870, 535)),
        ("AI-рекомендація", (570, 675)),
        ("Редагування профілю і фото", (900, 675)),
        ("Керування каталогом", (1040, 805)),
    ]
    centers = []
    for text, (cx, cy) in cases:
        box = (cx - 175, cy - 42, cx + 175, cy + 42)
        d.ellipse(box, fill="#E8F1FF", outline="#0A84FF", width=3)
        lines = textwrap.wrap(text, width=24)
        total_h = len(lines) * 25
        for i, line in enumerate(lines):
            tw = d.textlength(line, font=small_f)
            d.text((cx - tw / 2, cy - total_h / 2 + i * 26), line, font=small_f, fill="#0F172A")
        centers.append((text, cx, cy))

    def connect(actor_xy, case_xy, color="#64748B"):
        ax, ay = actor_xy
        cx, cy = case_xy
        d.line((ax, ay + 30, cx, cy), fill=color, width=3)

    for _, cx, cy in centers[:8]:
        connect((120, 320), (cx - 175, cy))
    for _, cx, cy in centers[6:9]:
        connect((1360, 650), (cx + 175, cy), "#7C3AED")
    connect((1360, 320), (1040 + 175, 805), "#0F766E")

    img.save(USE_CASE)


def make_code_shot(src: Path, start: int, end: int, out: Path, title: str) -> None:
    lines = src.read_text(encoding="utf-8").splitlines()[start - 1 : end]
    w, line_h, pad = 1500, 30, 34
    h = pad * 2 + 70 + line_h * len(lines)
    img = Image.new("RGB", (w, h), "#0B1020")
    d = ImageDraw.Draw(img)
    title_font = font(30, bold=True)
    code_font = font(22, mono=True)
    rounded(d, (18, 18, w - 18, h - 18), 26, "#111827", "#334155", 2)
    d.text((42, 34), title, font=title_font, fill="#E5E7EB")
    d.text((42, 72), str(src.relative_to(ROOT)), font=font(18), fill="#94A3B8")
    y = 116
    for idx, line in enumerate(lines, start=start):
        num = f"{idx:>4}"
        d.text((42, y), num, font=code_font, fill="#64748B")
        code = line.expandtabs(2)
        color = "#E5E7EB"
        stripped = code.strip()
        if stripped.startswith(("@app", "class ", "def ", "Future", "void ", "CREATE", "INSERT")):
            color = "#93C5FD"
        elif "required" in stripped or "return" in stripped:
            color = "#C4B5FD"
        elif stripped.startswith(("#", "--")):
            color = "#94A3B8"
        d.text((115, y), code[:105], font=code_font, fill=color)
        y += line_h
    img.save(out)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def style_cell(cell, bold=False, color="111827"):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.name = "Calibri"
            r.font.size = Pt(10)
            r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(color)


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Calibri"
        run.font.color.rgb = RGBColor(46, 116, 181)
    return p


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_image(doc: Document, path: Path, caption: str, width_in: float = 6.2) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(caption).bold = True
    p.paragraph_format.space_after = Pt(4)
    doc.add_picture(str(path), width=Inches(width_in))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def build_docx() -> None:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    make_use_case_diagram()
    make_code_shot(ROOT / "lib" / "main.dart", 1622, 1712, SCREEN_DIR / "07_flutter_ai_code.png", "Flutter: AI Agent Studio")
    make_code_shot(ROOT / "backend" / "main.py", 249, 333, SCREEN_DIR / "08_fastapi_code.png", "FastAPI: ratings, comments, recommendations")
    make_code_shot(ROOT / "backend" / "schema.sql", 1, 90, SCREEN_DIR / "09_postgresql_schema.png", "PostgreSQL: schema, comments and seed data")

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    for level in ["Heading 1", "Heading 2", "Heading 3"]:
        styles[level].font.name = "Calibri"
        styles[level].font.color.rgb = RGBColor(46, 116, 181)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Звіт за тиждень 5")
    run.font.name = "Calibri"
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("CineMind: персональні рекомендації фільмів з AI Agent Studio").italic = True

    meta = doc.add_table(rows=5, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.autofit = False
    labels = [
        ("ПІБ студента", "Іванов В.Є."),
        ("Команда", "Іванов В.Є."),
        ("Група", "КН-24007б"),
        ("Тема", "CineMind - мобільний застосунок для персональних рекомендацій фільмів"),
        ("GitHub", "https://github.com/Vanil-CEO/CineMind-movie-ai"),
    ]
    for row, (k, v) in zip(meta.rows, labels):
        row.cells[0].text = k
        row.cells[1].text = v
        set_cell_shading(row.cells[0], "E8EEF5")
        style_cell(row.cells[0], bold=True)
        style_cell(row.cells[1])

    add_heading(doc, "Діаграма прецедентів", 1)
    add_image(doc, USE_CASE, "Рисунок 1 - діаграма прецедентів", 6.4)

    add_heading(doc, "Мета тижня 5", 1)
    doc.add_paragraph(
        "Метою тижня 5 було суттєво відрізнити проєкт від стану 4 тижня: "
        "перейти від базового прототипу з оновленими афішами та backend-підготовкою "
        "до більш завершеного продуктового інтерфейсу з персоналізацією, AI Agent Studio, "
        "коментарями, оцінками, профілем користувача та production web-збіркою."
    )

    add_heading(doc, "Пророблена робота після 4 тижня", 1)
    add_bullets(
        doc,
        [
            "Flutter: створено cinematic home screen з великим hero-блоком, постерами, AI match та responsive-компонентами.",
            "Flutter: додано AI Agent Studio з пам'яттю агента, запитом користувача та персональними добірками.",
            "Flutter: додано редагування профілю, bio, фото через URL, компактну статистику та вплив профілю на рекомендації.",
            "Flutter: реалізовано коментарі до фільмів, анімоване оцінювання зірочками та оновлення персонального профілю після дій користувача.",
            "Flutter Web: зібрано production build і запущено статичну версію через локальний web-server.",
            "FastAPI: розширено backend API для каталогу фільмів, оцінок, списків, рекомендацій та AI-відповіді.",
            "PostgreSQL: підготовлено схему для users, movies, genres, actors, favorites, watch_later, watched_movies, user_ratings та seed-дані.",
            "Python: згенеровано стильні постери та підготовлено зображення для звітності.",
            "GitHub: проєкт прив'язано до репозиторію CineMind-movie-ai.",
        ],
    )

    add_heading(doc, "Порівняння з 4 тижнем", 1)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    hdr[0].text = "Напрям"
    hdr[1].text = "Стан після 4 тижня"
    hdr[2].text = "Результат 5 тижня"
    for cell in hdr:
        set_cell_shading(cell, "E8EEF5")
        style_cell(cell, bold=True)
    rows = [
        ("Інтерфейс", "Базовий Netflix-style прототип", "Cinematic iPhone/Apple TV style з великим hero-блоком"),
        ("Афіші", "Повернуті/оновлені афіші", "Уніфіковані poster cards, rails, responsive layout"),
        ("AI", "Демонстраційна відповідь", "AI Agent Studio з пам'яттю профілю, оцінками й коментарями"),
        ("Профіль", "Базовий профіль", "Редагування імені, bio, фото, статистика та AI memory"),
        ("Backend", "FastAPI/PostgreSQL налаштування", "Розширені моделі, endpoints, schema.sql, seed data"),
        ("Web", "Debug запуск", "Production build + локальний static web-server"),
    ]
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
            style_cell(cell)

    add_heading(doc, "Виконання функціональних вимог", 1)
    add_bullets(
        doc,
        [
            "Реєстрація/вхід користувача.",
            "Вибір і переналаштування кіновподобань.",
            "Персональна стрічка фільмів із AI match.",
            "Перегляд детальної інформації про фільм.",
            "Відкриття трейлерів через зовнішні посилання.",
            "Оцінювання фільмів зірочками.",
            "Коментування фільмів.",
            "Редагування профілю та фото.",
            "AI Agent Studio для персоналізованого підбору.",
            "FastAPI + PostgreSQL backend для зберігання та обробки даних.",
        ],
    )

    add_heading(doc, "Скріншоти застосунку", 1)
    app_images = [
        ("01_auth.png", "Рисунок 2 - екран реєстрації та входу"),
        ("02_cinematic_home.png", "Рисунок 3 - cinematic home screen"),
        ("03_tech_pipeline_posters.png", "Рисунок 4 - технологічний ланцюг та афіші"),
        ("04_ai_agent_studio.png", "Рисунок 5 - AI Agent Studio"),
        ("05_profile_memory.png", "Рисунок 6 - профіль користувача"),
        ("06_movie_details.png", "Рисунок 7 - детальна сторінка фільму"),
    ]
    for name, caption in app_images:
        add_image(doc, SCREEN_DIR / name, caption, 6.3)

    add_heading(doc, "Фрагменти реалізації", 1)
    code_images = [
        ("07_flutter_ai_code.png", "Рисунок 8 - Flutter-код AI Agent Studio"),
        ("08_fastapi_code.png", "Рисунок 9 - FastAPI endpoints"),
        ("09_postgresql_schema.png", "Рисунок 10 - PostgreSQL schema"),
    ]
    for name, caption in code_images:
        add_image(doc, SCREEN_DIR / name, caption, 6.3)

    add_heading(doc, "Використані технології", 1)
    add_bullets(
        doc,
        [
            "Flutter",
            "Dart",
            "Material 3",
            "FastAPI",
            "PostgreSQL",
            "psycopg",
            "Docker Compose",
            "Python",
            "python-docx",
            "Pillow",
            "GitHub",
        ],
    )

    add_heading(doc, "Перевірка", 1)
    add_bullets(
        doc,
        [
            "dart format",
            "flutter analyze",
            "flutter test",
            "flutter build web --release",
            "Локальний production web-server: http://127.0.0.1:5260/?v=ai-field-fix",
        ],
    )

    add_heading(doc, "Висновок", 1)
    doc.add_paragraph(
        "За тиждень 5 CineMind отримав помітний продуктовий розвиток після 4 тижня. "
        "Було реалізовано масштабний редизайн, персоналізацію, AI Agent Studio, оцінки, "
        "коментарі, профіль з фото, FastAPI endpoints, PostgreSQL schema та production web-збірку. "
        "Проєкт став більш завершеним і демонструє видимий прогрес у функціональності, "
        "архітектурі та користувацькому досвіді."
    )

    doc.save(OUT_DOCX)


if __name__ == "__main__":
    build_docx()
    print(OUT_DOCX)
